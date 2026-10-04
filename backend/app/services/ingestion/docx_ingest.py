"""
Word Document (.docx) Ingestion Service.
Extracts headings, paragraphs, bullet lists, and tables into ContentBlocks and ExtractedTables.
"""

import logging
from pathlib import Path
from typing import List, Tuple
import docx

from app.schemas.universal import ContentBlock, ExtractedTable, PageMeta

logger = logging.getLogger(__name__)


def ingest_docx_file(
    file_path: Path,
) -> Tuple[List[PageMeta], List[ContentBlock], List[ExtractedTable]]:
    """Ingests a .docx Word document."""
    doc = docx.Document(str(file_path))
    blocks: List[ContentBlock] = []
    tables: List[ExtractedTable] = []

    order = 0

    # 1. Paragraphs
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue

        style_name = p.style.name.lower() if p.style else ""
        block_type = "paragraph"
        if "heading" in style_name or "title" in style_name:
            block_type = "heading"
        elif "list" in style_name or text.startswith(("- ", "• ", "* ")):
            block_type = "list"

        blocks.append(
            ContentBlock(
                id=f"blk_docx_{order}",
                page=1,
                type=block_type,
                text=text,
                reading_order=order,
                ocr_conf=1.0,
            )
        )
        order += 1

    # 2. Tables
    for t_idx, tbl in enumerate(doc.tables):
        if not tbl.rows:
            continue
        headers = [cell.text.strip() for cell in tbl.rows[0].cells]
        rows = []
        for row in tbl.rows[1:]:
            rows.append([cell.text.strip() for cell in row.cells])

        tables.append(
            ExtractedTable(
                id=f"tbl_docx_{t_idx}",
                page=1,
                headers=headers,
                rows=rows,
                ocr_conf=1.0,
            )
        )

    pages = [
        PageMeta(
            n=1,
            width=612.0,
            height=792.0,
            source="docx",
        )
    ]

    return pages, blocks, tables
